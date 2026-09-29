# -*- coding: utf-8 -*-
"""Build DREAM CLASSES blackboard-style 16:9 PPTX for Class 10 Physics Ch.2."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

# ---------- palette ----------
BG      = (16, 22, 19)     # #101613
BG_TOP  = (20, 28, 25)     # gradient top
BG_BOT  = (11, 17, 14)     # gradient bottom
CARD    = (26, 35, 32)     # #1A2320
RED     = (200, 16, 46)    # #C8102E
GOLD    = (255, 193, 7)    # #FFC107
WHITE   = (255, 255, 255)
GREY    = (154, 165, 160)  # #9AA5A0
DARK    = (16, 22, 19)
HEADER  = (11, 16, 14)
DIVIDER = (42, 53, 49)

LATIN = "Segoe UI"
CSFONT = "Noto Sans Devanagari"

COACHING = "DREAM CLASSES KOTHWARA"
TEACHER = "Dream Sir"
MOBILE = "+91 97089 83294"
SUBJECT_HEAD = "PHYSICS  •  विद्युत"
CHAPTER_EN = "Electricity"
CHAPTER_HI = "विद्युत"

QUESTIONS = [
    ("विद्युत आवेश का SI मात्रक है —",
     ["वोल्ट", "कूलॉम", "ऐम्पियर", "ओम"]),
    ("विद्युत धारा का SI मात्रक है —",
     ["कूलॉम", "वोल्ट", "ऐम्पियर", "ओम"]),
    ("विभवांतर का SI मात्रक है —",
     ["ऐम्पियर", "वोल्ट", "ओम", "वाट"]),
    ("विद्युत प्रतिरोध का SI मात्रक है —",
     ["वोल्ट", "ऐम्पियर", "ओम", "वाट"]),
    ("ओम के नियम का सही सूत्र है —",
     ["V = IR", "I = VR", "R = IV", "V = I/R"]),
    ("ओम के नियम की खोज किसने की?",
     ["फैराडे", "ओम", "ऐम्पियर", "वोल्टा"]),
    ("प्रतिरोधकता का SI मात्रक है —",
     ["ओम", "ओम-मीटर", "वोल्ट", "ऐम्पियर"]),
    ("विद्युत शक्ति का SI मात्रक है —",
     ["जूल", "वाट", "वोल्ट", "कूलॉम"]),
    ("विद्युत ऊर्जा का व्यावहारिक मात्रक है —",
     ["वाट", "जूल", "किलोवाट-घंटा", "हॉर्स पावर"]),
    ("1 किलोवाट-घंटा बराबर होता है —",
     ["3.6 × 10⁶ जूल", "3.6 × 10⁵ जूल", "360 जूल", "1000 जूल"]),
    ("विद्युत धारा मापने के लिए प्रयुक्त यंत्र है —",
     ["वोल्टमीटर", "अमीटर", "ओममीटर", "वाटमीटर"]),
    ("विभवांतर मापा जाता है —",
     ["अमीटर से", "वोल्टमीटर से", "ओममीटर से", "वाटमीटर से"]),
    ("अमीटर को परिपथ में जोड़ा जाता है —",
     ["समांतर क्रम में", "श्रेणी क्रम में", "किसी भी क्रम में", "नहीं जोड़ा जाता"]),
    ("वोल्टमीटर को परिपथ में जोड़ा जाता है —",
     ["श्रेणी क्रम में", "समांतर क्रम में", "किसी भी क्रम में", "नहीं जोड़ा जाता"]),
    ("श्रेणी क्रम में जुड़े प्रतिरोधों का तुल्य प्रतिरोध होता है —",
     ["R = R₁ + R₂ + R₃", "1/R = 1/R₁ + 1/R₂ + 1/R₃", "R = R₁ × R₂ × R₃", "R = R₁ - R₂ - R₃"]),
    ("समांतर क्रम में जुड़े प्रतिरोधों का तुल्य प्रतिरोध होता है —",
     ["R = R₁ + R₂", "1/R = 1/R₁ + 1/R₂", "R = R₁ × R₂", "R = R₁ - R₂"]),
    ("जूल के ऊष्मीय नियम का सूत्र है —",
     ["H = I²Rt", "H = IR²t", "H = IRt²", "H = I²R/t"]),
    ("विद्युत शक्ति का सूत्र है —",
     ["P = VI", "P = V/I", "P = I/V", "P = V + I"]),
    ("बिजली के बल्ब का तंतु किस धातु का बना होता है?",
     ["ताँबा", "लोहा", "टंगस्टन", "ऐलुमिनियम"]),
    ("किसी चालक का प्रतिरोध निर्भर करता है —",
     ["लंबाई पर", "अनुप्रस्थ काट के क्षेत्रफल पर", "पदार्थ की प्रकृति पर", "उपर्युक्त सभी पर"]),
]


LABELS = ["(A)", "(B)", "(C)", "(D)"]


def _ensure_typeface(rPr, tag, typeface):
    el = rPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        el.set("typeface", typeface)
        rPr.append(el)
    else:
        el.set("typeface", typeface)


def style_run(run, size, bold=False, italic=False, color=WHITE,
              latin=LATIN, cs=CSFONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor(*color)
    f.name = latin
    rPr = run._r.get_or_add_rPr()
    _ensure_typeface(rPr, "a:ea", latin)
    _ensure_typeface(rPr, "a:cs", cs)
    rPr.set("szCs", str(int(size * 100)))
    rPr.set("bCs", "1" if bold else "0")
    rPr.set("iCs", "1" if italic else "0")


def add_run(p, text, **kw):
    r = p.add_run()
    r.text = text
    style_run(r, **kw)
    return r


def new_para(tf, first=False, align=PP_ALIGN.LEFT, ls=1.0, before=0, after=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = ls
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    return p


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP, ml=0.05, mr=0.05, mt=0.02, mb=0.02):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(ml)
    tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt)
    tf.margin_bottom = Inches(mb)
    return tf


def rect_shape(slide, l, t, w, h, fill=None, line=None, line_w=1.0,
               shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor(*fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.fill.solid()
        sp.line.fill.fore_color.rgb = RGBColor(*line)
        sp.line.width = Pt(line_w)
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def paint_bg(slide):
    try:
        fill = slide.background.fill
        fill.gradient()
        fill.gradient_angle = 90.0
        stops = fill.gradient_stops
        stops[0].color.rgb = RGBColor(*BG_TOP)
        stops[0].position = 0.0
        stops[1].color.rgb = RGBColor(*BG_BOT)
        stops[1].position = 1.0
    except Exception:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*BG)


def q_size(text):
    n = len(text)
    if n <= 55:
        return 46
    if n <= 80:
        return 42
    if n <= 110:
        return 38
    if n <= 150:
        return 34
    return 32


def opt_size(opts):
    m = max(len(o) for o in opts)
    if m <= 16:
        return 30
    if m <= 26:
        return 28
    if m <= 40:
        return 25
    if m <= 60:
        return 23
    return 21


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "VVI Objective Questions - Electricity - Dream Classes Kothwara"
prs.core_properties.author = "Dream Classes Kothwara"

# ================= TITLE SLIDE =================
s = prs.slides.add_slide(BLANK)
paint_bg(s)
rect_shape(s, 0, 0, 13.333, 0.06, fill=GOLD)  # top gold hairline

tf = textbox(s, 0.8, 1.05, 11.73, 0.75)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "विज्ञान (PHYSICS) — 10वीं कक्षा", size=32, bold=True, color=GOLD)

tf = textbox(s, 0.5, 1.85, 12.33, 1.5)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "VVI OBJECTIVE QUESTIONS", size=60, bold=True, color=WHITE)

rect_shape(s, 6.02, 3.5, 1.3, 0.045, fill=GOLD)

banner = rect_shape(s, 1.5, 3.8, 10.33, 1.2, fill=RED,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
btf = banner.text_frame
btf.word_wrap = True
btf.auto_size = None
btf.vertical_anchor = MSO_ANCHOR.MIDDLE
btf.margin_left = Inches(0.3); btf.margin_right = Inches(0.3)
p1 = new_para(btf, first=True, align=PP_ALIGN.CENTER, ls=1.0, after=4)
add_run(p1, CHAPTER_EN, size=28, bold=True, color=WHITE)
p2 = new_para(btf, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p2, CHAPTER_HI, size=26, bold=True, color=GOLD)

tf = textbox(s, 0.8, 5.45, 11.73, 1.3)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0, after=6)
add_run(p, COACHING, size=30, bold=True, color=GOLD)
p = new_para(tf, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, f"By: {TEACHER}   |   Mob: {MOBILE}", size=20, bold=False, color=WHITE)

# ================= QUESTION SLIDES =================
for idx, (qtext, opts) in enumerate(QUESTIONS, start=1):
    s = prs.slides.add_slide(BLANK)
    paint_bg(s)

    # header strip
    rect_shape(s, 0, 0, 13.333, 0.62, fill=HEADER)
    rect_shape(s, 0, 0.62, 13.333, 0.03, fill=GOLD)
    tf = textbox(s, 0.5, 0.06, 9.8, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.LEFT, ls=1.0)
    add_run(p, SUBJECT_HEAD, size=20, bold=True, color=WHITE)

    badge = rect_shape(s, 11.05, 0.10, 1.78, 0.42, fill=GOLD,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
    btf = badge.text_frame
    btf.word_wrap = True; btf.auto_size = None
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    btf.margin_left = Inches(0.05); btf.margin_right = Inches(0.05)
    p = new_para(btf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, "10TH CLASS", size=14, bold=True, color=DARK)

    # red banner
    banner = rect_shape(s, 0.5, 0.85, 12.333, 0.62, fill=RED,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    btf = banner.text_frame
    btf.word_wrap = True; btf.auto_size = None
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = new_para(btf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, "VVI OBJECTIVE QUESTION", size=24, bold=True, color=WHITE)
    rect_shape(s, 0.5, 1.47, 12.333, 0.035, fill=GOLD)

    # question number dot
    dot = rect_shape(s, 0.78, 0.94, 0.44, 0.44, fill=GOLD, shape=MSO_SHAPE.OVAL)
    dtf = dot.text_frame
    dtf.word_wrap = True; dtf.auto_size = None
    dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    dtf.margin_left = Inches(0); dtf.margin_right = Inches(0)
    p = new_para(dtf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, str(idx), size=20, bold=True, color=DARK)

    # question
    full_q = f"Q.{idx}. {qtext}"
    tf = textbox(s, 0.6, 1.68, 12.13, 2.0)
    p = new_para(tf, first=True, align=PP_ALIGN.LEFT, ls=1.06)
    add_run(p, full_q, size=q_size(full_q), bold=True, color=WHITE)

    # options 2x2
    osz = opt_size(opts)
    xs = [0.6, 6.79]
    ys = [3.88, 5.35]
    cw, ch = 5.94, 1.25
    for k, opt in enumerate(opts):
        card = rect_shape(s, xs[k % 2], ys[k // 2], cw, ch, fill=CARD,
                          line=GOLD, line_w=1.25,
                          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
        ctf = card.text_frame
        ctf.word_wrap = True; ctf.auto_size = None
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ctf.margin_left = Inches(0.22); ctf.margin_right = Inches(0.18)
        p = new_para(ctf, first=True, align=PP_ALIGN.LEFT, ls=1.04)
        add_run(p, LABELS[k] + "  ", size=osz, bold=True, color=GOLD)
        add_run(p, opt, size=osz, bold=False, color=WHITE)

    # footer
    rect_shape(s, 0.6, 6.82, 12.133, 0.015, fill=DIVIDER)
    tf = textbox(s, 0.6, 6.88, 5.5, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.LEFT, ls=1.0)
    add_run(p, "Dream Classes Kothwara", size=12, color=GREY)
    tf = textbox(s, 5.9, 6.88, 1.6, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, MOBILE, size=11, color=GREY)
    tf = textbox(s, 10.7, 6.88, 2.03, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.RIGHT, ls=1.0)
    add_run(p, TEACHER, size=12, color=GREY)

# ================= CLOSING SLIDE =================
s = prs.slides.add_slide(BLANK)
paint_bg(s)
rect_shape(s, 0, 0, 13.333, 0.06, fill=GOLD)

tf = textbox(s, 0.5, 1.5, 12.33, 1.5)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "THANK YOU", size=72, bold=True, color=WHITE)

tf = textbox(s, 0.5, 2.95, 12.33, 1.0)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "धन्यवाद!", size=44, bold=True, color=GOLD)

rect_shape(s, 6.02, 4.05, 1.3, 0.045, fill=GOLD)

tf = textbox(s, 0.8, 4.4, 11.73, 1.3)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0, after=6)
add_run(p, COACHING, size=30, bold=True, color=GOLD)
p = new_para(tf, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, f"By: {TEACHER}   |   Mob: {MOBILE}", size=20, color=WHITE)

tf = textbox(s, 0.8, 5.95, 11.73, 0.6)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "Physics  •  विद्युत  •  10वीं कक्षा", size=16, color=GREY)

out = "DREAMCLASSES_PHYSICS_Electricity.pptx"
prs.save(out)
print(f"saved {out}: {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
