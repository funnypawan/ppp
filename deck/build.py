# -*- coding: utf-8 -*-
"""
DREAM CLASSES KOTHWARA — जैव विज्ञान अध्याय 1 (जैव जगत) : 100 वस्तुनिष्ठ प्रश्न, 1 प्रश्न प्रति स्लाइड।
हर स्लाइड के ऊपर "DREAM CLASSES KOTHWARA" और ठीक नीचे "BY – DREAM SIR"।
आउटपुट : ../Biology_Ch1_Jaiv_Jagat_100_MCQ_Dream_Classes.pptx
"""
import os, math, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

from questions import Q

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.abspath(os.path.join(HERE, "..", "assets"))
OUT = os.path.abspath(os.path.join(HERE, "..", "Biology_Ch1_Jaiv_Jagat_100_MCQ_Dream_Classes.pptx"))

SW, SH = 13.333, 7.5
HI = "Nirmala UI"
SCHOOL = "DREAM CLASSES KOTHWARA"
BY = "BY – DREAM SIR"
CH = "अध्याय 1 : जैव जगत  (The Living World)"
BOARD = "जैव विज्ञान • कक्षा 11 (इंटरमीडिएट) • BSEB वस्तुनिष्ठ अभ्यास"
LETTERS = ["A", "B", "C", "D"]

# ---------------------------------------------------------------- palette
PALETTES = [
    ("FFF6E5", "FFE3B8", "E85D04", "DC2F02", ("FFE8A3", "CDEBFF", "D6F5D0", "FBD5E8"), "2B2118"),
    ("EAF6FF", "CBE9FF", "0353A4", "0466C8", ("D7ECFF", "FFE9A8", "D8F8E5", "FBD3E3"), "102A43"),
    ("F1FFF2", "D2F5D8", "1B7F5A", "10694E", ("DDF7D6", "FFF0BF", "D9EEFF", "F9D5E7"), "12352A"),
    ("FFF0F5", "FBD3E4", "C9184A", "A4133C", ("FFDDE7", "FFF3C0", "D7F0FF", "E3D7FF"), "4A1226"),
    ("F6F3FF", "E0D8FF", "5A189A", "7B2CBF", ("E9DFFF", "FFE6B8", "D4F5EA", "FFD9E2"), "241344"),
    ("FFFBEA", "FFEFC2", "9A6700", "BC6C25", ("FFF1B8", "D3EFFB", "E7DFFF", "FFDCC7"), "3D2C00"),
    ("E9FBFF", "C7F0FA", "0B7285", "087F8C", ("D2F4FB", "FFE9B0", "DEFFD9", "F5D6F0"), "073B4C"),
    ("FFF1EC", "FFD9C7", "BA1200", "9D0208", ("FFE0D6", "D9EEFF", "E1F8D8", "F3DEFB"), "3D0800"),
]
OK_FILL, OK_LINE, OK_INK = "C7F5D6", "0F8A3E", "0B5A2A"

# ---------------------------------------------------------------- primitives
def _in(v):
    return Inches(v)

def rect(slide, x, y, w, h, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16):
    s = slide.shapes.add_shape(shape, _in(x), _in(y), _in(max(w, 0.05)), _in(max(h, 0.05)))
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    return s

def grad(shape, stops, angle=90.0):
    spPr = shape._element.spPr
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    x = '<a:gradFill %s gradRotateWithShape="1"><a:gsLst>' % nsdecls('a')
    for pos, hx in stops:
        x += '<a:gs pos="%d"><a:srgbClr val="%s"/></a:gs>' % (int(pos * 100000), hx)
    x += '</a:gsLst><a:lin ang="%d" scaled="1"/></a:gradFill>' % int(angle * 60000)
    spPr.insert_element_before(parse_xml(x), 'a:ln', 'a:effectLst', 'a:effectDag',
                               'a:scene3d', 'a:sp3d', 'a:extLst')

def solid(shape, hx, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hx)
    if alpha is not None:
        spPr = shape._element.spPr
        sf = spPr.find(qn('a:solidFill'))
        clr = sf.find(qn('a:srgbClr')) if sf is not None else None
        if clr is not None:
            clr.append(parse_xml('<a:alpha %s val="%d"/>' % (nsdecls('a'), max(2000, min(100000, int(alpha))))))

def no_line(shape):
    shape.line.fill.background()

def line(shape, hx, w=1.4, dash=None):
    shape.line.color.rgb = RGBColor.from_string(hx)
    shape.line.width = Pt(w)
    if dash:
        shape.line.dash_style = dash

def soft_shadow(shape, blur=0.07, dist=0.028, alpha=40000):
    spPr = shape._element.spPr
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)
    spPr.append(parse_xml(
        '<a:effectLst %s><a:outerShdw blurRad="%d" dist="%d" dir="5400000" rotWithShape="0">'
        '<a:srgbClr val="141432"><a:alpha val="%d"/></a:srgbClr></a:outerShdw></a:effectLst>'
        % (nsdecls('a'), Emu(_in(blur)), Emu(_in(dist)), alpha)))

def script_fonts(run):
    run.font.name = HI
    rPr = run._r.get_or_add_rPr()
    rPr.set('lang', 'hi-IN')
    rPr.set('altLang', 'en-US')
    latin = rPr.find(qn('a:latin'))
    for tag in ('a:ea', 'a:cs'):
        if rPr.find(qn(tag)) is None:
            latin.addnext(parse_xml('<%s %s typeface="%s"/>' % (tag, nsdecls('a'), HI)))
    ea, cs = rPr.find(qn('a:ea')), rPr.find(qn('a:cs'))
    if ea is not None:
        latin.addnext(ea)
    if cs is not None and ea is not None:
        ea.addnext(cs)

def _fill_para(p, runs, size, bold, color, italic, align, line_spacing):
    """write runs into paragraph; '\n' becomes a real <a:br/> (PowerPoint-safe)"""
    p.alignment = align
    p.line_spacing = line_spacing
    last = None
    for txt, ov in runs:
        for j, part in enumerate(str(txt).split('\n')):
            if j:
                br = parse_xml('<a:br %s/>' % nsdecls('a'))
                if last is not None:
                    _rPr = last._r.find(qn('a:rPr'))
                    if _rPr is not None:
                        br.append(copy.deepcopy(_rPr))      # keep the same line metrics
                p._p.append(br)
            if part == '':
                continue
            r = p.add_run()
            r.text = part
            r.font.size = Pt(ov.get('size', size))
            r.font.bold = ov.get('bold', bold)
            r.font.italic = ov.get('italic', italic)
            r.font.color.rgb = RGBColor.from_string(ov.get('color', color))
            script_fonts(r)
            last = r

def text(shape, runs, size=16, bold=False, color="10233F", align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE, italic=False, line_spacing=1.0, margins=(0.12, 0.10, 0.05, 0.05)):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    l, r_, t, b = margins
    tf.margin_left, tf.margin_right = _in(l), _in(r_)
    tf.margin_top, tf.margin_bottom = _in(t), _in(b)
    if isinstance(runs, str):
        runs = [(runs, {})]
    _fill_para(tf.paragraphs[0], runs, size, bold, color, italic, align, line_spacing)
    return tf

def more_para(shape, runs, size=13, bold=False, color="33415C", align=PP_ALIGN.CENTER,
              space_before=2, line_spacing=1.0, italic=False):
    tf = shape.text_frame
    if isinstance(runs, str):
        runs = [(runs, {})]
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    _fill_para(p, runs, size, bold, color, italic, align, line_spacing)
    return p

def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
    tb.text_frame.word_wrap = True
    tb.text_frame.margin_left = tb.text_frame.margin_right = _in(0)
    tb.text_frame.margin_top = tb.text_frame.margin_bottom = _in(0)
    tb.shadow.inherit = False
    return tb

def picture(slide, key, x, y, max_w, max_h):
    path = os.path.join(ASSETS, key + ".jpg")
    if not os.path.exists(path):
        return None
    from PIL import Image
    iw, ih = Image.open(path).size
    ar = iw / ih
    w, h = max_w, max_w / ar
    if h > max_h:
        h, w = max_h, max_h * ar
    return slide.shapes.add_picture(path, _in(x + (max_w - w) / 2), _in(y), width=_in(w), height=_in(h))

# ---------------------------------------------------------------- text metrics (approx, conservative)
def _chw(ch, latin_scale=1.0):
    o = ord(ch)
    if ch == ' ':
        return 0.27
    if 0x0900 <= o <= 0x097F:          # Devanagari
        if ch in '्  ँ ं ः ':
            return 0.30
        return 0.545 * latin_scale
    if ch.isdigit():
        return 0.53
    if ch.isupper():
        return 0.63
    if ch in 'iljft.,:;\'"!()[]|/':
        return 0.29
    if ch in 'mw':
        return 0.72
    if ch in '-—–_':
        return 0.42
    if 0x2190 <= o <= 0x2BFF or o in (0x2713, 0x2714, 0x2705):   # symbols / ticks
        return 0.62
    if o > 0x1F000:                                               # emoji
        return 1.05
    return 0.50

def est_width_in(s, size_pt, bold=False):
    f = (1.05 if bold else 1.0) * size_pt / 72.0
    return sum(_chw(c) for c in s) * f

def est_lines(s, size_pt, width_in, bold=False):
    """greedy word-wrap line count"""
    words, lines, cur = s.split(' '), 1, 0.0
    avail = max(0.4, width_in)
    for wd in words:
        w = est_width_in(wd + ' ', size_pt, bold)
        if cur + w > avail and cur > 0:
            lines += 1
            cur = w
        else:
            cur += w
    return lines

def fit_height(lines, size_pt, spacing=1.2):
    return lines * size_pt * spacing / 72.0

def choose_size(text_, width_in, sizes, max_h, bold=False, spacing=1.2):
    for s in sizes:
        h = fit_height(est_lines(text_, s, width_in, bold), s, spacing)
        if h <= max_h:
            return s, est_lines(text_, s, width_in, bold), h
    s = sizes[-1]
    return s, est_lines(text_, s, width_in, bold), fit_height(est_lines(text_, s, width_in, bold), s, spacing)

# ---------------------------------------------------------------- decorations
def decor(slide, pal):
    A, B = pal[2], pal[3]
    for (cx, cy, d, al) in ((11.75, -0.85, 3.2, 15000), (-1.05, 5.15, 3.5, 12000), (11.55, 6.05, 2.15, 11000)):
        s = rect(slide, cx, cy, d, d, MSO_SHAPE.OVAL)
        solid(s, B, alpha=al)
        no_line(s)
    for i, (bx, by, bs) in enumerate(((0.40, 1.92, 0.12), (12.80, 1.92, 0.12), (0.40, 6.62, 0.10), (12.82, 6.62, 0.10))):
        s = rect(slide, bx, by, bs, bs, MSO_SHAPE.OVAL)
        solid(s, A, alpha=55000 if i % 2 else 33000)
        no_line(s)

def header(slide, pal):
    b = rect(slide, 0.30, 0.20, SW - 0.60, 1.02, radius=0.22)
    grad(b, [(0.0, pal[2]), (1.0, pal[3])], 32)
    no_line(b)
    soft_shadow(b, 0.09, 0.035, 36000)
    text(b, [(SCHOOL, {'size': 27, 'bold': True, 'color': 'FFFFFF'})], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.TOP, margins=(0.3, 0.3, 0.08, 0))
    more_para(b, [(BY, {'size': 14.5, 'bold': True, 'color': 'FFE79A'})], align=PP_ALIGN.CENTER, space_before=1)
    st = rect(slide, 0.30, 1.265, SW - 0.60, 0.055, MSO_SHAPE.RECTANGLE, radius=0.5)
    grad(st, [(0.0, 'FFD166'), (0.34, '06D6A0'), (0.68, '118AB2'), (1.0, 'EF476F')], 0)
    no_line(st)

def subbar(slide, pal, num, total=100):
    s = rect(slide, 0.30, 1.40, SW - 0.60, 0.42, radius=0.5)
    solid(s, 'FFFFFF', alpha=62000)
    line(s, pal[2], 1.0)
    text(s, [(CH, {'size': 13, 'bold': True, 'color': pal[2]}),
             ("     |     " + BOARD, {'size': 11, 'color': '3B4A63'})],
         align=PP_ALIGN.LEFT, margins=(0.22, 2.5, 0.0, 0.0))
    bd = rect(slide, SW - 2.30, 1.43, 1.98, 0.36, radius=0.5)
    solid(bd, pal[3])
    no_line(bd)
    text(bd, "प्रश्न %d / %d" % (num, total), size=12.5, bold=True, color='FFFFFF',
         align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0.0, 0.0))

# ---------------------------------------------------------------- mini diagrams (scaled + auto-fit captions)
def diagram(kind, slide, pal, X, Y, W, H):
    A, B, INK, T = pal[2], pal[3], pal[5], pal[4]
    card = rect(slide, X, Y, W, H, radius=0.07)
    solid(card, 'FFFFFF', alpha=82000)
    line(card, A, 1.5, MSO_LINE.DASH)

    S = min(1.0, H / 4.10, W / 3.65)
    # graphic (non-stacked) kinds keep their natural proportions but get centred in a tall panel
    NAT = {"growth": 3.28, "cell": 2.95, "family": 2.90, "herbarium": 2.95,
           "garden": 3.55, "museum": 2.95, "key": 3.00}.get(kind)
    TO = max(0.0, (H - NAT * min(1.0, W / 3.65)) / 2.0) if NAT else 0.0
    def d(v):
        return v * S
    def fs(v):
        return max(7.4, round(v * S, 1))

    def chip_box(x, y, w, h, fill, line_hx=None, lw=0.9, radius=0.30):
        b = rect(slide, X + x, Y + y, w, h, radius=radius)
        solid(b, fill)
        line(b, line_hx or 'FFFFFF', lw)
        return b

    def cap(y, txt, size, color, x=0.14, w=None, bottom=None, align=PP_ALIGN.CENTER, bold=True):
        """caption that shrinks itself so it never runs past the panel bottom"""
        w = (W - 2 * x) if w is None else w
        lim = (bottom if bottom is not None else H)
        segs = str(txt).split('\n')
        chosen, hh = size, 0.0
        for cand in (size, size - .7, size - 1.4, size - 2.1, size - 2.8, 8.2, 7.6):
            n = max(est_lines(sg, fs(cand), w - 0.12) for sg in segs)
            hh = (n * len(segs)) * fs(cand) * 1.26 / 72.0
            chosen = fs(cand)
            if y + hh <= lim or cand <= 7.6:
                break
        box = textbox(slide, X + x, Y + TO + y, w, hh + 0.06)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, sg in enumerate(segs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            _fill_para(p, [(sg, {})], size=chosen, bold=bold, color=color, italic=False,
                       align=align, line_spacing=1.0)
        return hh

    if kind == 'hierarchy':
        labels = ["जगत (Kingdom)", "Division / Phylum", "वर्ग (Class)", "क्रम (Order)",
                  "कुल (Family)", "गोत्र (Genus)", "जाति (Species)"]
        y = d(0.14)
        rh = max(d(0.325), (H - d(0.28) - d(0.60)) / 7.0)
        bh = min(rh - d(0.045), d(0.46))
        for i, lab in enumerate(labels):
            w = W - d(0.90) - i * d(0.155)
            x = d(0.45) - i * d(0.078)
            b = chip_box(x, y, w, bh, T[i % 4])
            text(b, lab, size=fs(10.6), bold=(i == 6), color=INK, align=PP_ALIGN.CENTER,
                 margins=(0.02, 0.02, 0, 0))
            y += rh
        cap(y + d(0.04), "↑ ऊपर जाते ही — साझा लक्षण घटते हैं; लक्षण अधिक सामान्य एवं मौलिक बनते हैं",
            9.6, A, bottom=H)

    elif kind == 'binomial':
        rows = [("Mangifera", "गोत्र • Capital अक्षर", T[1]),
                ("indica", "जाति-विशेषण • small अक्षर", T[2]),
                ("Linn.", "लेखक का संक्षिप्त नाम", T[3])]
        y = d(0.16)
        rh = max(d(0.50), (H - d(0.32) - d(1.05)) / 3.0)
        for val, lab, fl in rows:
            b = chip_box(d(0.24), y, W - d(0.48), min(rh - d(0.07), d(0.66)), fl, lw=1.0)
            text(b, [(val + "   ", {'size': fs(15), 'bold': True, 'italic': True, 'color': INK}),
                     (lab, {'size': fs(9.2), 'color': '3B4A63'})],
                 align=PP_ALIGN.CENTER, margins=(0.03, 0.03, 0, 0))
            y += rh
        cap(y + d(0.06), "मुद्रित → Italics\nहस्तलिखित → दोनों शब्द अलग-अलग underline\nएक ही टैक्सॉन के लिए एक ही नाम",
            9.8, A, bottom=H)

    elif kind == 'traits':
        items = [("वृद्धि (growth)", T[0]), ("प्रजनन (reproduction)", T[1]),
                 ("चयापचय (metabolism)", T[2]), ("चेतना (consciousness)", T[3])]
        y = d(0.14)
        rh = max(d(0.45), (H - d(0.28) - d(0.82)) / 4.0)
        for lab, fl in items:
            b = chip_box(d(0.22), y, W - d(0.44), min(rh - d(0.06), d(0.56)), fl)
            text(b, [("✓  ", {'size': fs(12), 'bold': True, 'color': OK_LINE}),
                     (lab, {'size': fs(11.4), 'bold': True, 'color': INK})],
                 align=PP_ALIGN.LEFT, margins=(0.12, 0.06, 0, 0))
            y += rh
        cap(y + d(0.02), "कोई एक लक्षण अकेला 'जीव' को परिभाषित नहीं करता —\nसभी लक्षण सामूहिक रूप से जीवन दर्शाते हैं",
            9.6, A, bottom=H)

    elif kind == 'growth':
        cap(d(0.10), "वृद्धि = आंतरिक कोशिका-विभाजन", 11.4, A, bold=True)
        for k, (ttl, fl, bars) in enumerate((("पादप : अपरिमित", T[2], (.20, .32, .44, .56, .68)),
                                             ("जंतु : सीमित", T[0], (.22, .42, .56, .57, .57)))):
            bw_ = (W - d(0.62)) / 2
            x0 = d(0.22) + k * (bw_ + d(0.18))
            box = rect(slide, X + x0, Y + TO + d(0.50), bw_, d(1.60), radius=0.08)
            solid(box, 'FFFFFF')
            line(box, fl, 1.2)
            bx = x0 + d(0.09)
            cw = (bw_ - d(0.20)) / 5 - d(0.04)
            for hgt in bars:
                bar = rect(slide, X + bx, Y + TO + d(0.50) + d(1.46) - d(hgt), cw, d(hgt), MSO_SHAPE.RECTANGLE)
                solid(bar, fl)
                no_line(bar)
                bx += cw + d(0.05)
            cap(d(2.14), ttl, 10.2, INK, x=x0, w=bw_)
        cap(d(2.50), "निर्जीव (पर्वत) — सतह पर जमाव से बढ़ता है;\nजीवों में वृद्धि भीतर से — अतः वृद्धि एकल निर्णायक लक्षण नहीं",
            9.4, B, bottom=H)

    elif kind == 'cell':
        cap(d(0.10), "कोशिकीय संगठन = जीवन का निर्णायक लक्षण", 10.4, A)
        for k, (lab, sub, dia, fl) in enumerate((("प्रागुकेंद्रिक", "DNA खुला", 1.00, T[1]),
                                                 ("सुकेंद्रिक", "प्रकेंद्रक युक्त", 1.22, T[2]))):
            col = d(0.26) + k * (W - d(0.52)) / 2
            dia_ = d(dia)
            cx = col + (W - d(0.52)) / 4 - dia_ / 2
            c = rect(slide, X + cx, Y + TO + d(0.52), dia_, dia_, MSO_SHAPE.OVAL)
            solid(c, fl)
            line(c, A, 1.4)
            if k:
                nd = dia_ * 0.36
                n2 = rect(slide, X + cx + (dia_ - nd) / 2, Y + TO + d(0.52) + (dia_ - nd) / 2, nd, nd, MSO_SHAPE.OVAL)
                solid(n2, B)
                no_line(n2)
            cap(d(0.52) + dia_ + d(0.04), lab + "\n" + sub, 9.8, INK, x=col, w=(W - d(0.52)) / 2,
                bottom=H - d(0.50))
        cap(d(2.30), "विषाणु — अकोशिकीय (non-cellular); इनका स्वर विशेष बहस का विषय", 9.6, B, bottom=H)

    elif kind == 'family':
        top = chip_box(d(0.28), d(0.18), W - d(0.56), d(0.40), A, lw=0.1, radius=0.22)
        text(top, "कुल : Solanaceae", size=fs(12.4), bold=True, color='FFFFFF', align=PP_ALIGN.CENTER,
             margins=(0.02, 0.02, 0, 0))
        names = [("Solanum", "आलू, बैंगन"), ("Petunia", "गुलदौरा"), ("Datura", "धतूरा")]
        cw = (W - d(0.60)) / 3
        for i, (g, sp) in enumerate(names):
            cx = d(0.22) + i * (cw + d(0.08))
            ln = rect(slide, X + cx + cw / 2 - d(0.012), Y + TO + d(0.58), d(0.024), d(0.20), MSO_SHAPE.RECTANGLE)
            solid(ln, A)
            no_line(ln)
            b = chip_box(cx, d(0.78), cw, d(0.74), T[i % 4], lw=1.0)
            text(b, [(g, {'size': fs(11.2), 'bold': True, 'italic': True, 'color': INK}),
                     (sp, {'size': fs(8.8), 'color': '44546C'})],
                 align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))
        cap(d(1.58), "अनेक संबंधित गोत्र → कुल\n(vegetative + reproductive दोनों लक्षण)", 9.6, A, bottom=H - d(0.62))
        cap(d(2.16), "क्रम : Convolvulaceae + Solanaceae → Polymoniales\n(आधार — पुष्प-लक्षण)", 9.4, B, bottom=H)

    elif kind == 'table':
        rows = [("कुल", "Hominidae / Poaceae"), ("क्रम", "Primata / Poales"),
                ("वर्ग", "Mammalia / Monocot."), ("संघ / Division", "Chordata / Angiospermae"),
                ("जाति", "sapiens / aestivum")]
        y = d(0.14)
        rh = max(d(0.41), (H - d(0.28) - d(0.72)) / 5.0)
        for j, (k, v) in enumerate(rows):
            a1 = chip_box(d(0.20), y, (W - d(0.40)) * 0.37, min(rh - d(0.06), d(0.50)), T[j % 4])
            text(a1, k, size=fs(9.0), bold=True, color=INK, align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))
            a2 = rect(slide, X + d(0.24) + (W - d(0.40)) * 0.37, Y + y, (W - d(0.40)) * 0.60, min(rh - d(0.06), d(0.50)), radius=0.3)
            solid(a2, 'FFFFFF')
            line(a2, A, 0.9)
            text(a2, v, size=fs(8.8), bold=True, color=A, align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))
            y += rh
        cap(y + d(0.03), "मनुष्य : Homo sapiens  •  गेहूँ : Triticum aestivum\n(Table 1.1 के अनुसार)", 9.2, B, bottom=H)

    elif kind == 'five':
        ks = [("जगत (Kingdom)", 'FFFFFF', A, True), ("Monera — जीवाणु", T[0], INK, False),
              ("Protista — एककोशिकीय", T[1], INK, False), ("Fungi — कवक", T[2], INK, False),
              ("Plantae — पादप", T[3], INK, False), ("Animalia — जंतु", T[0], INK, False)]
        y = d(0.12)
        gap5 = max(d(0.05), (H - d(0.24) - d(0.72)) / 6.0 - d(0.355))
        for lab, fl, ink, head in ks:
            h = min(d(0.42), d(0.32 if head else 0.355) + gap5)
            b = chip_box(d(0.24), y, W - d(0.48), h, fl, lw=0.9)
            text(b, lab, size=fs(11.4 if head else 10.6), bold=True, color=ink, align=PP_ALIGN.CENTER,
                 margins=(0.02, 0.02, 0, 0))
            y += h + gap5
        cap(y + d(0.03), "R.H. Whittaker (1969)\nआधार: कोशिका-प्रकार, कोशिका भित्ति, पोषण, शारीरिक संगठन, जातिवृत्ति",
            9.2, A, bottom=H)

    elif kind == 'domain':
        ds = [("Archaea", "अति उष्ण/लवणीय आवास"), ("Bacteria", "सच्चे जीवाणु (Eubacteria)"),
              ("Eukarya", "सुकेंद्रिक — पादप, जंतु, कवक")]
        y = d(0.16)
        rh = max(d(0.58), (H - d(0.32) - d(0.62)) / 3.0)
        for i, (lab, sub) in enumerate(ds):
            b = chip_box(d(0.24), y, W - d(0.48), min(rh - d(0.09), d(0.78)), T[i], lw=1.0)
            text(b, [(lab, {'size': fs(12.4), 'bold': True, 'color': INK}),
                     (("   " + sub), {'size': fs(8.8), 'color': '43536B'})],
                 align=PP_ALIGN.CENTER, margins=(0.03, 0.03, 0, 0))
            y += rh
        cap(y + d(0.04), "कार्ल वूज (1990) — rRNA/जातिवृत्तीय संबंधों पर आधारित तीन डोमेन", 9.4, A, bottom=H)

    elif kind == 'herbarium':
        sh = rect(slide, X + d(0.26), Y + TO + d(0.16), W - d(0.52), d(1.90), MSO_SHAPE.RECTANGLE, radius=0.03)
        solid(sh, 'FFFDF5')
        line(sh, A, 1.6)
        st = rect(slide, X + W / 2 - d(0.014), Y + TO + d(0.30), d(0.028), d(1.10), MSO_SHAPE.RECTANGLE)
        solid(st, '2D6A4F')
        no_line(st)
        for i, dy in enumerate((0.02, 0.28, 0.54, 0.80)):
            for sgn in (-1, 1):
                lf = rect(slide, X + W / 2 + (d(0.26) if sgn > 0 else -d(0.58)), Y + TO + d(0.36) + d(dy),
                          d(0.32), d(0.13), MSO_SHAPE.OVAL)
                solid(lf, '52B788')
                no_line(lf)
                lf.rotation = 16 * sgn
        fl = rect(slide, X + W / 2 - d(0.11), Y + TO + d(0.13), d(0.22), d(0.18), MSO_SHAPE.OVAL)
        solid(fl, 'EF476F')
        no_line(fl)
        for ty in (0.50, 1.02):
            tp = rect(slide, X + d(0.60), Y + TO + d(ty), W - d(1.20), d(0.07), MSO_SHAPE.RECTANGLE)
            solid(tp, 'D9CBA8', alpha=88000)
            no_line(tp)
        lb = rect(slide, X + d(0.40), Y + TO + d(1.56), W - d(0.80), d(0.44), radius=0.16)
        solid(lb, 'FFFFFF')
        line(lb, A, 1.0)
        cap(d(1.60), "Family • Genus • Species |\nस्थानीय नाम, उपयोगी भाग, तिथि, आवास", 8.2, INK,
            x=d(0.46), w=W - d(0.92), bottom=d(2.02))
        cap(d(2.14), "शीट ≈ 41 × 29 से.मी. — मान्य वर्गीकरण पद्धति के अनुसार रैक में", 9.2, A, bottom=H)

    elif kind == 'steps':
        st = [("1", "पादप-संग्रहण + field catalogue क्रमांक"), ("2", "पौधा-प्रेस में दबाकर सुखाना"),
              ("3", "शीट पर fix + कीट-नाशक"), ("4", "लेबल; रैक में वर्गीक्रम अनुसार")]
        y = d(0.14)
        rh = max(d(0.60), (H - d(0.28) - d(0.50)) / 4.0)
        for n, lab in st:
            cds = min(d(0.40), rh * 0.55)
            c = rect(slide, X + d(0.22), Y + y + (rh - d(0.56)) / 2 - d(0.06), cds, cds, MSO_SHAPE.OVAL)
            solid(c, B)
            no_line(c)
            text(c, n, size=fs(11.5), bold=True, color='FFFFFF', align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))
            cap(y - d(0.02), lab, 10.2, INK, x=d(0.62), w=W - d(0.84), bottom=y + rh + d(0.02),
                align=PP_ALIGN.LEFT)
            if n != "4":
                ar = rect(slide, X + d(0.36), Y + y + cds + d(0.03), d(0.045), max(d(0.10), rh - cds - d(0.10)), MSO_SHAPE.RECTANGLE)
                solid(ar, A)
                no_line(ar)
            y += rh
        cap(y + d(0.02), "हर्बेरियम = त्वरित संदर्भ (quick referral)", 9.4, A, bottom=H)

    elif kind == 'garden':
        dome = rect(slide, X + d(0.46), Y + TO + d(0.52), W - d(0.92), d(0.92), radius=0.30)
        solid(dome, 'CDEFFB')
        line(dome, A, 1.4)
        roof = rect(slide, X + d(0.62), Y + TO + d(0.16), W - d(1.24), d(0.48), MSO_SHAPE.ISOSCELES_TRIANGLE)
        solid(roof, '7FC8F8')
        no_line(roof)
        for i in range(3):
            p_ = rect(slide, X + d(0.70) + i * (W - d(1.60)) / 2.15, Y + TO + d(0.94), d(0.26), d(0.26), MSO_SHAPE.OVAL)
            solid(p_, ['52B788', 'FFB703', 'EF476F'][i])
            no_line(p_)
        for i, (lab, val) in enumerate((("क्यू (Kew)", "1759 • >45,000 जातियाँ"),
                                        ("शिबपुर, हावड़ा", "स्थापना 1787"),
                                        ("NBRI, लखनऊ", "प्रसिद्ध केन्द्र"))):
            ggap = max(d(0.42), (H - TO - d(2.00)) / 3.0)
            b = chip_box(d(0.22), d(1.56) + i * ggap, W - d(0.44), min(d(0.44), ggap - d(0.05)), T[i], lw=0.9)
            text(b, [(lab + "  ", {'size': fs(10.2), 'bold': True, 'color': INK}),
                     (val, {'size': fs(8.6), 'color': '43536B'})],
                 align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))
        cap(d(2.90), "वनस्पति उद्यान = जीवित पादप + वैज्ञानिक/सामान्य नाम का लेबल", 9.4, A, bottom=H)

    elif kind == 'museum':
        for i in range(3):
            jw = (W - d(0.68)) / 3
            jx = d(0.28) + i * (jw + d(0.05))
            jar = rect(slide, X + jx, Y + TO + d(0.26), jw, d(0.86), radius=0.22)
            solid(jar, 'EAF7FF')
            line(jar, A, 1.2)
            liq = rect(slide, X + jx + d(0.06), Y + TO + d(0.52), jw - d(0.12), d(0.48), MSO_SHAPE.RECTANGLE)
            solid(liq, '9AD8E8')
            no_line(liq)
            lid = rect(slide, X + jx + d(0.10), Y + TO + d(0.17), jw - d(0.20), d(0.13), MSO_SHAPE.RECTANGLE)
            solid(lid, B)
            no_line(lid)
        bx = rect(slide, X + d(0.26), Y + TO + d(1.26), W - d(0.52), d(0.64), radius=0.12)
        solid(bx, 'FFF3D6')
        line(bx, A, 1.2)
        for r in range(2):
            for c in range(4):
                bf = rect(slide, X + d(0.44) + c * d(0.58), Y + TO + d(1.34) + r * d(0.25), d(0.28), d(0.17), MSO_SHAPE.HEXAGON)
                solid(bf, ['FFB703', 'EF476F', '06D6A0', '118AB2'][(r + c) % 4])
                no_line(bf)
        cap(d(1.98), "द्रव : 70% एल्कोहॉल • बड़े प्राणी : stuffed\nकीट : insect boxes • उद्देश्य : अध्ययन एवं संदर्भ",
            9.2, A, bottom=H)

    elif kind == 'key':
        c1 = rect(slide, X + d(0.24), Y + TO + d(0.18), W - d(0.48), d(0.78), radius=0.16)
        solid(c1, T[1])
        line(c1, A, 1.0)
        text(c1, [("1a  ", {'size': fs(10.2), 'bold': True, 'color': B}),
                  ("एक लक्षण का कथन", {'size': fs(10.2), 'bold': True, 'color': INK}),
                  ("\n1b  ", {'size': fs(10.2), 'bold': True, 'color': B}),
                  ("विपरीत (opposing) कथन", {'size': fs(10.2), 'bold': True, 'color': INK})],
             align=PP_ALIGN.LEFT, margins=(0.14, 0.08, 0.02, 0.02))
        cap(d(1.02), "युग्मक (couplet) — एक स्वीकार = दूसरा अस्वीकार", 9.2, A, x=0.18, bold=True)
        b1 = chip_box(d(0.24), d(1.36), (W - d(0.56)) / 2 - d(0.04), d(0.56), T[2], lw=1.0)
        text(b1, "टैक्सॉन की पहचान", size=fs(9.6), bold=True, color=INK, align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))
        b2 = chip_box(d(0.32) + (W - d(0.56)) / 2, d(1.36), (W - d(0.56)) / 2 - d(0.04), d(0.56), T[3], lw=1.0)
        text(b2, "अगला युग्मक 2a / 2b", size=fs(9.4), bold=True, color=INK, align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))
        cap(d(2.02), "कुंजी विश्लेषणात्मक — वर्णन नहीं देती\nभारत में कुल-कुंजी : Bentham & Hooker", 9.4, B, bottom=H)

    elif kind == 'aids':
        items = [("हर्बेरियम", "शुष्क-पीसे पत्रक"), ("वनस्पति उद्यान", "जीवित पादप"),
                 ("संग्रहालय", "संरक्षित नमूने"), ("कुंजी", "विश्लेषणात्मक"),
                 ("साहित्य", "Flora • Manual • Monograph")]
        y = d(0.14)
        rh = max(d(0.47), (H - d(0.28) - d(0.50)) / 5.0)
        for i, (lab, sub) in enumerate(items):
            b = chip_box(d(0.22), y, W - d(0.44), min(rh - d(0.07), d(0.58)), T[i % 4], lw=0.9)
            text(b, [(lab, {'size': fs(11.2), 'bold': True, 'color': INK}),
                     (("   " + sub), {'size': fs(8.6), 'color': '43536B'})],
                 align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))
            y += rh
        cap(y + d(0.02), "वर्गीकरण-अध्ययन को सुगम बनाने वाले साधन", 9.4, A, bottom=H)
    return

# ---------------------------------------------------------------- question slide
def question_slide(prs, idx, rec, pal):
    qq, opts, ans, expl, vis, img, tag = rec
    s = prs.slides.add_slide(prs.slide_layouts[6])

    bg = rect(s, 0, 0, SW, SH, MSO_SHAPE.RECTANGLE)
    grad(bg, [(0.0, pal[0]), (0.58, pal[1]), (1.0, pal[0])], 118)
    no_line(bg)
    decor(s, pal)
    header(s, pal)
    subbar(s, pal, idx)

    has_side = bool(vis) or bool(img)
    X0, M = 0.30, 0.30
    body_w = (SW - 2 * M - 0.14 - 3.62) if has_side else (SW - 2 * M)
    qtext_w = body_w - 0.62

    # --- question card (height auto-fits the question text)
    qsizes = [21.5, 20.5, 19.5, 18.5, 17.5, 16.5, 15.5, 14.5]
    qsize, qlines, qh_text = choose_size(qq, qtext_w, qsizes, 1.02, bold=True, spacing=1.24)
    chip_h = 0.36
    card_h = chip_h + qh_text + 0.26
    cy = 1.94
    card = rect(s, X0, cy, body_w, card_h, radius=0.09)
    solid(card, 'FFFFFF')
    line(card, pal[2], 1.5)
    soft_shadow(card)
    bar = rect(s, X0, cy, 0.10, card_h, MSO_SHAPE.RECTANGLE)
    grad(bar, [(0.0, pal[2]), (1.0, pal[3])], 90)
    no_line(bar)

    nb = rect(s, X0 + 0.24, cy + 0.075, 0.98, chip_h, radius=0.5)
    solid(nb, pal[3])
    no_line(nb)
    text(nb, "प्र. %d" % idx, size=12.5, bold=True, color='FFFFFF', align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))

    tw = est_width_in(tag, 11, True) + 0.34
    tagc = rect(s, X0 + body_w - tw - 0.16, cy + 0.075, tw, chip_h, radius=0.5)
    solid(tagc, pal[4][idx % 4])
    no_line(tagc)
    text(tagc, tag, size=11, bold=True, color='10233F', align=PP_ALIGN.CENTER, margins=(0.02, 0.02, 0, 0))

    qt = textbox(s, X0 + 0.31, cy + chip_h + 0.11, qtext_w, qh_text + 0.10)
    text(qt, qq, size=qsize, bold=True, color=pal[5], align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0), line_spacing=0.98)

    # --- options
    oy = cy + card_h + 0.14
    obot = 6.38
    gap = 0.085
    rowh = (obot - oy - gap * 3) / 4.0
    ow = body_w - 1.52
    osizes = [17.5, 16.5, 15.5, 14.5, 13.5, 12.5]
    osize = osizes[0]
    for cand in osizes:
        if max(fit_height(est_lines(o, cand, ow, i == ans), cand, 1.24) for i, o in enumerate(opts)) <= rowh - 0.14:
            osize = cand
            break
    else:
        osize = osizes[-1]
    for i, op in enumerate(opts):
        y = oy + i * (rowh + gap)
        correct = (i == ans)
        row = rect(s, X0, y, body_w, rowh, radius=0.20)
        if correct:
            solid(row, OK_FILL)
            line(row, OK_LINE, 2.1)
            soft_shadow(row, 0.055, 0.022, 28000)
        else:
            solid(row, 'FFFFFF', alpha=80000)
            line(row, pal[2], 1.0)
        cd = 0.42
        c = rect(s, X0 + 0.13, y + (rowh - cd) / 2, cd, cd, MSO_SHAPE.OVAL)
        solid(c, OK_LINE if correct else pal[4][i])
        no_line(c)
        text(c, LETTERS[i], size=13.5, bold=True, color='FFFFFF' if correct else pal[5],
             align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))
        ot = textbox(s, X0 + 0.66, y + 0.03, ow, rowh - 0.06)
        text(ot, op, size=osize, bold=correct, color=OK_INK if correct else pal[5],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0), line_spacing=0.96)
        if correct:
            tk = rect(s, X0 + body_w - 0.56, y + (rowh - 0.32) / 2, 0.40, 0.32, MSO_SHAPE.OVAL)
            solid(tk, 'FFFFFF')
            line(tk, OK_LINE, 1.3)
            text(tk, "✓", size=13.5, bold=True, color=OK_LINE, align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))

    # --- side panel : image + diagram
    if has_side:
        px = X0 + body_w + 0.14
        pw = SW - M - px
        py, ph = 1.94, obot - 1.94
        panel = rect(s, px, py, pw, ph, radius=0.06)
        solid(panel, 'FFFFFF', alpha=45000)
        no_line(panel)
        pad = 0.10
        top = py + pad
        if img:
            ih = 1.30 if vis else ph - 2 * pad
            picture(s, img, px + pad, top, pw - 2 * pad, ih)
            top += ih + 0.10
            if not vis:
                cap = textbox(s, px + pad + 0.05, top, pw - 2 * pad - 0.10, ph - (top - py) - pad)
                text(cap, tag, size=12.5, bold=True, color=pal[2], align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0))
        if vis:
            diagram(vis, s, pal, px + 0.05, top, pw - 0.10, py + ph - top - pad + 0.05)

    # --- answer strip (label pill + explanation, both auto-fitted)
    ay, ah = 6.50, 0.62
    aw_full = (SW - 2 * M) if not has_side else body_w
    aS = rect(s, X0, ay, aw_full, ah, radius=0.26)
    grad(aS, [(0.0, '0F8A3E'), (1.0, '05B45F')], 18)
    no_line(aS)
    soft_shadow(aS, 0.05, 0.02, 26000)
    lbl = "✅  उत्तर : (%s)" % LETTERS[ans]
    lsz = 14.0
    lw_ = est_width_in(lbl, lsz, True) + 0.34
    ex_w = aw_full - lw_ - 0.42
    esize = 12.0
    for cand in (12.0, 11.4, 10.8, 10.2, 9.6, 9.0, 8.4):
        if fit_height(est_lines(expl, cand, ex_w), cand, 1.26) <= ah - 0.14:
            esize = cand
            break
    else:
        esize = 8.4
    lb2 = textbox(s, X0 + 0.16, ay + 0.04, lw_ - 0.06, ah - 0.08)
    text(lb2, lbl, size=lsz, bold=True, color='FFFFFF', align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
    ex = textbox(s, X0 + lw_ + 0.14, ay + 0.04, ex_w, ah - 0.08)
    text(ex, expl, size=esize, bold=False, color='E9FFF2', align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0), line_spacing=1.0)
    if has_side:
        pb = rect(s, X0 + body_w + 0.14, ay, SW - M - (X0 + body_w + 0.14), ah, radius=0.26)
        grad(pb, [(0.0, pal[2]), (1.0, pal[3])], 18)
        no_line(pb)
        text(pb, "%d / 100" % idx, size=13, bold=True, color='FFFFFF', align=PP_ALIGN.CENTER, margins=(0, 0, 0, 0))

    # --- teacher notes (speaker notes) : answer + reason + tag
    try:
        nt = s.notes_slide.notes_text_frame
        nt.text = ("प्रश्न %d — %s\n%s\n\n%s\n\nउत्तर : (%s) %s\nकारण : %s"
                   % (idx, tag, qq, "\n".join("%s. %s" % (LETTERS[k], o) for k, o in enumerate(opts)),
                      LETTERS[ans], opts[ans], expl))
        for p_ in nt.paragraphs:
            for r_ in p_.runs:
                script_fonts(r_)
    except Exception:
        pass

    # --- footer
    ft = textbox(s, X0, 7.14, SW - 2 * M, 0.28)
    text(ft, [(SCHOOL + "  •  " + BY, {'size': 9, 'bold': True, 'color': pal[2]}),
              ("      जैव विज्ञान • अध्याय 1 जैव जगत • 100 वस्तुनिष्ठ प्रश्न (BSEB कक्षा 11)", {'size': 9, 'color': '55647E'})],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
    return s

# ---------------------------------------------------------------- title slide
def title_slide(prs, pal):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(s, 0, 0, SW, SH, MSO_SHAPE.RECTANGLE)
    grad(bg, [(0.0, pal[0]), (0.42, pal[1]), (1.0, 'FFFFFF')], 126)
    no_line(bg)
    decor(s, pal)

    band = rect(s, 0.30, 0.24, SW - 0.60, 1.10, radius=0.20)
    grad(band, [(0.0, pal[2]), (1.0, pal[3])], 28)
    no_line(band)
    soft_shadow(band, 0.10, 0.04, 38000)
    text(band, [(SCHOOL, {'size': 31, 'bold': True, 'color': 'FFFFFF'})], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.TOP, margins=(0.3, 0.3, 0.07, 0))
    more_para(band, [(BY, {'size': 15.5, 'bold': True, 'color': 'FFE79A'})], align=PP_ALIGN.CENTER)

    picture(s, "hero", 7.60, 1.70, 5.30, 3.60)

    t1 = textbox(s, 0.60, 1.72, 6.80, 1.46)
    text(t1, [("जैव विज्ञान  •  कक्षा 11", {'size': 19, 'bold': True, 'color': pal[3]}),
              ("\nअध्याय 1 : ", {'size': 25, 'bold': True, 'color': '17324F'}),
              ("जैव जगत", {'size': 38, 'bold': True, 'color': pal[2]}),
              ("   The Living World", {'size': 17, 'bold': True, 'italic': True, 'color': '3B5A78'})],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0), line_spacing=1.0)

    t2 = textbox(s, 0.60, 3.22, 6.80, 0.66)
    text(t2, [("100 ", {'size': 32, 'bold': True, 'color': 'D62828'}),
              ("वस्तुनिष्ठ प्रश्न ", {'size': 22, 'bold': True, 'color': '17324F'}),
              ("(BSEB Board Level)", {'size': 17, 'bold': True, 'color': pal[2]})],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))

    chips = ["BSEB कक्षा 11 • अध्याय-आधारित", "प्रति स्लाइड 1 प्रश्न + 4 विकल्प",
             "उत्तर + संक्षिप्त कारण सहित", "NCERT तथ्यों पर पूर्णतः आधारित",
             "छोटे आरेख, चित्र व टैग", "अंत में सम्पूर्ण उत्तर-कुंजी"]
    x, y = 0.60, 3.98
    for i, lab in enumerate(chips):
        w = 3.28
        c = rect(s, x, y, w, 0.44, radius=0.42)
        solid(c, 'FFFFFF')
        line(c, pal[2], 1.3)
        dot = rect(s, x + 0.12, y + 0.135, 0.17, 0.17, MSO_SHAPE.OVAL)
        solid(dot, pal[4][i % 4])
        no_line(dot)
        text(c, "  " + lab, size=12, bold=True, color='17324F', align=PP_ALIGN.LEFT, margins=(0.26, 0.06, 0, 0))
        if i % 2:
            y += 0.55
            x = 0.60
        else:
            x = 0.60 + 3.44

    note = rect(s, 0.60, 5.74, 6.60, 1.02, radius=0.12)
    grad(note, [(0.0, pal[2]), (1.0, pal[3])], 24)
    no_line(note)
    soft_shadow(note)
    text(note, [("उपयोग विधि :  ", {'size': 13, 'bold': True, 'color': 'FFE79A'}),
                ("प्रथम विकल्प स्वयं चुनिए, तत्पश्चात् नीचे दिए ", {'size': 12.5, 'color': 'FFFFFF'}),
                ("✅ उत्तर", {'size': 12.5, 'bold': True, 'color': 'FFF3B0'}),
                (" व कारण से मिलाइए। हर 25 प्रश्न पर एक बार रिवीज़न ज़रूर करें।", {'size': 12.5, 'color': 'FFFFFF'})],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.20, 0.16, 0.04, 0.04), line_spacing=1.04)

    t3 = textbox(s, 7.55, 5.55, 5.35, 1.30)
    text(t3, [("100 प्रश्न • 100 स्लाइड • उत्तर-कुंजी सहित", {'size': 12.5, 'bold': True, 'color': pal[2]}),
              ("\nतैयार : DREAM CLASSES KOTHWARA\n" + BY, {'size': 12, 'bold': True, 'color': '37475E'})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0), line_spacing=1.1)
    return s

# ---------------------------------------------------------------- answer key slides
def key_slide(prs, pal, start, end, part):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(s, 0, 0, SW, SH, MSO_SHAPE.RECTANGLE)
    grad(bg, [(0.0, pal[0]), (0.5, pal[1]), (1.0, pal[0])], 118)
    no_line(bg)
    decor(s, pal)
    head = rect(s, 0.30, 0.22, SW - 0.60, 0.94, radius=0.24)
    grad(head, [(0.0, pal[2]), (1.0, pal[3])], 30)
    no_line(head)
    text(head, [(SCHOOL, {'size': 22, 'bold': True, 'color': 'FFFFFF'})], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.TOP, margins=(0.2, 0.2, 0.05, 0))
    more_para(head, [(BY + "      |      उत्तर-कुंजी (Answer Key) — भाग %d : प्रश्न %d – %d" % (part, start, end),
                      {'size': 12.5, 'bold': True, 'color': 'FFE79A'})], align=PP_ALIGN.CENTER)
    cols, rows = 5, 10
    cw = (SW - 0.60 - 0.10 * (cols - 1)) / cols
    chh = 0.478
    y0 = 1.32
    for n in range(start, end + 1):
        i = n - start
        cx = 0.30 + (i % cols) * (cw + 0.10)
        cy = y0 + (i // cols) * (chh + 0.05)
        qq, opts, ans, expl, vis, img, tag = Q[n - 1]
        card = rect(s, cx, cy, cw, chh, radius=0.22)
        solid(card, 'FFFFFF', alpha=86000)
        line(card, pal[2], 0.9)
        opt = opts[ans]
        avail = cw - 0.15 - est_width_in("प्र %d  →  (%s) " % (n, LETTERS[ans]), 12.5, True)
        lsz, osz, txt = 9.4, 9.4, opt
        while txt and est_width_in(txt, osz) > avail:
            cut = max(6, int(len(txt) * 0.85))
            if cut == len(txt):
                lsz = osz = max(7.4, osz - 0.6)
                if osz <= 7.5 and est_width_in(txt, osz) <= avail:
                    break
            txt = txt[:cut].rstrip(' ,;:') + "…"
        text(card, [("प्र %d" % n, {'size': 10, 'bold': True, 'color': '4A5A73'}),
                     ("  →  ", {'size': 10, 'color': '9AA7BB'}),
                     ("(%s) " % LETTERS[ans], {'size': 12.5, 'bold': True, 'color': OK_LINE}),
                     (txt, {'size': osz, 'color': pal[5]})],
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.10, 0.05, 0, 0), line_spacing=0.9)
    ft = textbox(s, 0.30, 6.66, SW - 0.60, 0.60)
    text(ft, [("टिप्पणी : ", {'size': 11, 'bold': True, 'color': pal[2]}),
              ("प्रश्न-स्लाइड पर हर उत्तर के साथ संक्षिप्त कारण भी दिया गया है; यह केवल त्वरित जाँच हेतु कुंजी है।",
               {'size': 11, 'color': '3B4A63'})],
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.04, 0.04, 0, 0))
    return s

def closing_slide(prs, pal):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(s, 0, 0, SW, SH, MSO_SHAPE.RECTANGLE)
    grad(bg, [(0.0, pal[2]), (0.55, pal[3]), (1.0, '14233C')], 122)
    no_line(bg)
    for (cx, cy, dd, al) in ((10.75, -0.95, 3.7, 20000), (-1.15, 4.60, 3.9, 16000), (5.85, 5.45, 2.7, 13000)):
        o = rect(s, cx, cy, dd, dd, MSO_SHAPE.OVAL)
        solid(o, 'FFFFFF', alpha=al)
        no_line(o)
    t = textbox(s, 0.9, 1.42, SW - 1.8, 1.70)
    text(t, [("अभ्यास जारी रखिए!", {'size': 43, 'bold': True, 'color': 'FFFFFF'}),
             ("\n100 प्रश्न पूरे — अब उत्तर-कुंजी से self-check कीजिए।", {'size': 19, 'bold': True, 'color': 'FFE79A'})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0), line_spacing=1.05)
    card = rect(s, 2.90, 3.42, SW - 5.80, 1.92, radius=0.10)
    solid(card, 'FFFFFF', alpha=92000)
    no_line(card)
    soft_shadow(card, 0.10, 0.05, 40000)
    text(card, [(SCHOOL, {'size': 27, 'bold': True, 'color': pal[2]}),
                ("\n" + BY, {'size': 17, 'bold': True, 'color': 'D62828'}),
                ("\nजैव विज्ञान • अध्याय 1 जैव जगत • BSEB वस्तुनिष्ठ श्रृंखला", {'size': 13, 'color': '43536B'})],
         align=PP_ALIGN.CENTER, margins=(0.2, 0.2, 0.1, 0.1))
    b = textbox(s, 0.9, 5.72, SW - 1.8, 0.9)
    text(b, [("अगला सेट : अध्याय 2 — जैविक वर्गीकरण (Biological Classification) • 100 वस्तुनिष्ठ प्रश्न",
              {'size': 14, 'bold': True, 'color': 'FFFFFF'})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0))
    return s

# ---------------------------------------------------------------- build
def build(path=OUT):
    prs = Presentation()
    prs.slide_width, prs.slide_height = _in(SW), _in(SH)
    title_slide(prs, PALETTES[0])
    for i, rec in enumerate(Q, 1):
        question_slide(prs, i, rec, PALETTES[(i - 1) % len(PALETTES)])
    key_slide(prs, PALETTES[2], 1, 50, 1)
    key_slide(prs, PALETTES[6], 51, 100, 2)
    closing_slide(prs, PALETTES[4])
    cp = prs.core_properties
    cp.title = "जैव जगत (The Living World) — 100 वस्तुनिष्ठ प्रश्न | DREAM CLASSES KOTHWARA"
    cp.author = "DREAM CLASSES KOTHWARA (By – Dream Sir)"
    cp.subject = "जैव विज्ञान अध्याय 1 • BSEB कक्षा 11 • वस्तुनिष्ठ"
    cp.keywords = "BSEB, Class 11, Biology, The Living World, Objective, Hindi, DREAM CLASSES KOTHWARA"
    cp.comments = ("100 objective questions, one per slide, with 4 options in Hindi, "
                   "answer + reason, colourful design and small diagrams. Facts as per NCERT Class 11 Ch-1.")
    prs.save(path)
    return path, len(prs.slides._sldIdLst)

if __name__ == "__main__":
    p, n = build()
    print("saved:", p, "| slides:", n)
