# wireframe preview renderer (geometry + fills + text-blocks) for layout sanity checking
import sys, os
from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw
EMU = 914400.0
SCALE = 110  # px per inch

def inch(v): return (v or 0) / EMU

def hexof(spPr, default=None):
    for tag in ('a:solidFill', 'a:gradFill'):
        f = spPr.find(qn(tag))
        if f is None:
            continue
        srgb = f.find('.//' + qn('a:srgbClr'))
        if srgb is not None:
            a = srgb.find(qn('a:alpha'))
            return srgb.get('val'), (int(a.get('val')) / 100000.0 if a is not None else 1.0)
        stops = f.findall('.//' + qn('a:gs'))
        if stops:
            c = stops[0].find(qn('a:srgbClr'))
            if c is not None:
                return c.get('val'), 1.0
    return default

def walk(shapes):
    for sh in shapes:
        yield sh

def render(pptx, idx_list, out):
    prs = Presentation(pptx)
    W, H = int(inch(prs.slide_width) * SCALE), int(inch(prs.slide_height) * SCALE)
    tiles = []
    for si in idx_list:
        sl = list(prs.slides)[si - 1]
        img = Image.new('RGB', (W, H), (255, 255, 255))
        d = ImageDraw.Draw(img, 'RGBA')
        for sh in walk(sl.shapes):
            x, y = inch(sh.left) * SCALE, inch(sh.top) * SCALE
            w, h = inch(sh.width) * SCALE, inch(sh.height) * SCALE
            st = getattr(sh, 'shape_type', None)
            if st is not None and str(st).startswith('PICTURE'):
                try:
                    blob = sh.image.blob
                    from io import BytesIO
                    pic = Image.open(BytesIO(blob)).convert('RGB')
                    pic.thumbnail((max(1, int(w)), max(1, int(h))))
                    img.paste(pic, (int(x + (w - pic.width) / 2), int(y + (h - pic.height) / 2)))
                except Exception:
                    pass
                continue
            spPr = sh._element.find(qn('p:spPr')) if hasattr(sh._element, 'find') else None
            fill = None
            if spPr is not None:
                r = hexof(spPr)
                if r:
                    hx, al = r
                    rgb = tuple(int(hx[i:i + 2], 16) for i in (0, 2, 4))
                    fill = rgb + (int(255 * (0.35 + 0.65 * al)),)
            geom = sh._element.find(qn('p:spPr'))
            prst = None
            if geom is not None:
                pg = geom.find(qn('a:prstGeom'))
                if pg is not None:
                    prst = pg.get('prst')
            col = fill or (245, 245, 245, 120)
            if prst == 'ellipse':
                d.ellipse([x, y, x + w, y + h], fill=col, outline=(60, 60, 60, 120))
            elif prst == 'rect':
                d.rectangle([x, y, x + w, y + h], fill=col, outline=(60, 60, 60, 90))
            else:
                rad = max(2, min(w, h) * 0.18)
                d.rounded_rectangle([x, y, x + w, y + h], radius=rad, fill=col, outline=(60, 60, 60, 110))
            # text blocks drawn as bars so crowding is visible
            if sh.has_text_frame and sh.text_frame.text.strip():
                ty = y + h * 0.18
                for p in sh.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    if not txt:
                        continue
                    size = max([(r.font.size.pt if r.font.size else 16) for r in p.runs])
                    tw = sum(0.5 * size * (1.05 if r.font.bold else 1.0) / 72.0 * SCALE * max(1, len(r.text))
                             for r in p.runs) / 1.0
                    nlines = max(1, int(tw / max(10, w - 12)) + 1)
                    for k in range(nlines):
                        bh = size * 1.35 / 72.0 * SCALE * 0.62
                        d.rectangle([x + 5, ty + k * (bh + 2), x + w - 5, ty + k * (bh + 2) + bh],
                                    fill=(30, 30, 40, 105))
                    ty += nlines * (size * 1.35 / 72.0 * SCALE * 0.62 + 2) + 3
        d.rectangle([0, 0, W - 1, H - 1], outline=(200, 60, 60), width=3)
        d.text((8, 6), "slide %d" % si, fill=(200, 20, 20))
        tiles.append(img)
    cols = 2
    rows = (len(tiles) + 1) // 2
    sheet = Image.new('RGB', (cols * W + (cols + 1) * 12, rows * H + (rows + 1) * 12), (190, 190, 195))
    for i, t in enumerate(tiles):
        sheet.paste(t, (12 + (i % cols) * (W + 12), 12 + (i // cols) * (H + 12)))
    sheet.save(out, quality=88)
    print("saved", out, sheet.size)

if __name__ == '__main__':
    render(os.path.join('..', 'Biology_Ch1_Jaiv_Jagat_100_MCQ_Dream_Classes.pptx'),
           [int(a) for a in sys.argv[1:-1]] if len(sys.argv) > 2 else [1, 5], sys.argv[-1])
