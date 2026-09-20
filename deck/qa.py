# -*- coding: utf-8 -*-
"""QA जाँच — बनी PPTX को पढ़कर layout/flow/उत्तर-संगति की जाँच करता है।"""
import copy, math, os, sys
from pptx import Presentation
from pptx.oxml.ns import qn
from questions import Q

EMU = 914400.0
def inch(v):
    return (v or 0) / EMU

def _chw(ch):
    o = ord(ch)
    if ch == ' ':
        return .27
    if 0x0900 <= o <= 0x097F:
        return .30 if ch in '् ँ ं ः' else .545
    if ch.isdigit():
        return .53
    if ch.isupper():
        return .63
    if ch in 'iljft.,:;\'"!()[]|/':
        return .29
    if ch in 'mw':
        return .72
    if o > 0x1F000:
        return 1.05
    return .50

def est_w(txt, size, bold=False):
    return sum(_chw(c) for c in txt) * size * (1.05 if bold else 1.0) / 72.0

def para_lines(p, width):
    """count rendered lines honouring <a:br/> and word wrap"""
    segs, cur, size, bold = [], "", 18, False
    for child in p._p:
        tag = child.tag.split('}')[1]
        if tag == 'br':
            segs.append((cur, size, bold)); cur = ""
        elif tag == 'r':
            t = child.find(qn('a:t'))
            rPr = child.find(qn('a:rPr'))
            sz = rPr.find(qn('a:latin')) if rPr is not None else None
            if rPr is not None and rPr.get('sz'):
                size = int(rPr.get('sz')) / 100.0
            if rPr is not None:
                bold = rPr.get('b') == '1'
            txt = (t.text or '') if t is not None else ''
            for word in txt.split(' '):
                w = est_w(word + ' ', size, bold)
                if cur and est_w(cur + ' ', size, bold) + w > width:
                    segs.append((cur, size, bold)); cur = word
                else:
                    cur = (cur + ' ' + word).strip() if cur else word
    if cur:
        segs.append((cur, size, bold))
    ls = p.line_spacing or 1.0
    if isinstance(ls, float):
        return max(1, len(segs)), ls
    return max(1, len(segs)), 1.0

def box_height(p, width):
    n, ls = para_lines(p, width)
    mx = 0.0
    for child in p._p:
        if child.tag.endswith('}r') or child.tag.endswith('}br'):
            rPr = child.find(qn('a:rPr'))
            if rPr is not None and rPr.get('sz'):
                mx = max(mx, int(rPr.get('sz')) / 100.0)
    mx = mx or 10.0
    return n * mx * 1.24 * ls / 72.0

def check(path):
    prs = Presentation(path)
    SW, SH = inch(prs.slide_width), inch(prs.slide_height)
    problems, notes = [], []
    n_slides = len(prs.slides._sldIdLst)
    for si, sl in enumerate(prs.slides, 1):
        texts = []
        boxes = []
        for sh in sl.shapes:
            x, y, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text
                texts.append(t)
                ml, mr = inch(sh.text_frame.margin_left), inch(sh.text_frame.margin_right)
                mt, mb = inch(sh.text_frame.margin_top), inch(sh.text_frame.margin_bottom)
                need = sum(box_height(p, max(.3, w - ml - mr)) for p in sh.text_frame.paragraphs) + mt + mb
                if need > h + 0.05:
                    problems.append((si, 'TEXT-OVERFLOW need %.2f > box %.2f' % (need, h), t[:44].replace('\n', ' / ')))
                if x < -0.02 or y < -0.02 or x + w > SW + 0.02 or y + h > SH + 0.02:
                    problems.append((si, 'TEXT-OUT-OF-SLIDE', t[:44].replace('\n', ' / ')))
                boxes.append((x, y, w, h, t[:28]))
            if sh.shape_type == 13 and (x < -0.02 or y < -0.02 or x + w > SW + 0.02 or y + h > SH + 0.02):
                problems.append((si, 'PICTURE-OUT-OF-SLIDE', 'picture'))
        blob = " ".join(texts)
        if "DREAM CLASSES KOTHWARA" not in blob:
            problems.append((si, 'MISSING-HEADER', 'DREAM CLASSES KOTHWARA absent'))
        if "DREAM SIR" not in blob:
            problems.append((si, 'MISSING-BYLINE', 'BY – DREAM SIR absent'))
        if 2 <= si <= 101:      # question slides
            for need_txt in ("उत्तर : (", "प्र. "):
                if need_txt not in blob:
                    problems.append((si, 'MISSING-' + need_txt, ''))
            ok = sum(1 for s_ in sl.shapes if s_.has_text_frame and s_.text_frame.text.startswith("✓"))
            notes.append(si)
    # answer integrity: correct option text must appear on its slide
    for i, (qq, opts, ans, expl, vis, img, tag) in enumerate(Q, 1):
        sl = list(prs.slides)[i]                     # +1 because slide 1 is the title
        blob = " ".join(s_.text_frame.text for s_ in sl.shapes if s_.has_text_frame)
        if opts[ans][:24] not in blob:
            problems.append((i + 1, 'ANSWER-TEXT-MISSING', opts[ans][:24]))
        if qq[:24] not in blob:
            problems.append((i + 1, 'QUESTION-TEXT-MISSING', qq[:24]))
    return n_slides, problems, notes

if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else os.path.join('..', 'Biology_Ch1_Jaiv_Jagat_100_MCQ_Dream_Classes.pptx')
    n, problems, notes = check(path)
    print("slides:", n, "| question slides checked:", len(notes))
    if problems:
        print("PROBLEMS (%d):" % len(problems))
        for p_ in problems[:40]:
            print("  ", p_)
    else:
        print("कोई समस्या नहीं — सभी स्लाइड साफ़ ✔")
