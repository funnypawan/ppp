# -*- coding: utf-8 -*-
"""Build DREAM CLASSES blackboard-style 16:9 PPTX for Class 10 Physics Ch.2."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

# ---------- palette ----------
BG      = (16, 22, 19)     # #101613
BG_TOP  = (20, 28, 25)     # gradient top
BG_BOT  = (11, 17, 14)     # gradient bottom
CARD    = (26, 35, 32)     # #1A2320
RED     = (200, 16, 46)    # #C8102E
GOLD    = (255, 193, 7)    # #FFC107
WHITE   = (255, 255, 255)
GREY    = (154, 165, 160)  # #9AA5A0
DARK    = (16, 22, 19)
HEADER  = (11, 16, 14)
DIVIDER = (42, 53, 49)

LATIN = "Segoe UI"
CSFONT = "Noto Sans Devanagari"

COACHING = "DREAM CLASSES KOTHWARA"
TEACHER = "Dream Sir"
MOBILE = "+91 97089 83294"
SUBJECT_HEAD = "PHYSICS  •  प्रकाश का अपवर्तन"

QUESTIONS = [
    ("प्रकाश की चाल सबसे अधिक किस माध्यम में होती है?",
     ["काँच में", "पानी में", "निर्वात में", "हीरे में"]),
    ("स्नेल का नियम किससे संबंधित है?",
     ["प्रकाश के परावर्तन से", "प्रकाश के अपवर्तन से", "प्रकाश के विक्षेपण से", "प्रकाश के ध्रुवण से"]),
    ("अपवर्तनांक (n) का सही सूत्र कौन-सा है?",
     ["n = sin i / sin r", "n = sin r / sin i", "n = v / c", "n = sin i × sin r"]),
    ("प्रकाश की किरण विरल माध्यम से सघन माध्यम में प्रवेश करती है, तो वह —",
     ["अभिलंब की ओर झुकती है", "अभिलंब से दूर हटती है", "बिना झुके सीधी जाती है", "परावर्तित हो जाती है"]),
    ("प्रकाश की किरण सघन माध्यम से विरल माध्यम में जाती है, तो वह —",
     ["अभिलंब की ओर झुकती है", "अभिलंब से दूर हटती है", "बिना झुके सीधी जाती है", "अवशोषित हो जाती है"]),
    ("वायु के सापेक्ष काँच का अपवर्तनांक लगभग कितना होता है?",
     ["1.00", "1.33", "1.50", "2.42"]),
    ("वायु के सापेक्ष जल का अपवर्तनांक कितना होता है?",
     ["1.00", "1.33", "1.50", "2.42"]),
    ("लेंस की क्षमता का SI मात्रक क्या है?",
     ["मीटर", "डायोप्टर", "वाट", "हर्ट्ज"]),
    ("उत्तल लेंस की फोकस दूरी होती है —",
     ["धनात्मक", "ऋणात्मक", "शून्य", "अनंत"]),
    ("अवतल लेंस की फोकस दूरी होती है —",
     ["धनात्मक", "ऋणात्मक", "शून्य", "अनंत"]),
    ("तारों का टिमटिमाना किस घटना के कारण होता है?",
     ["परावर्तन", "अपवर्तन", "विक्षेपण", "प्रकीर्णन"]),
    ("पानी में आंशिक रूप से डूबी छड़ मुड़ी हुई प्रतीत होती है —",
     ["परावर्तन के कारण", "अपवर्तन के कारण", "विक्षेपण के कारण", "विवर्तन के कारण"]),
    ("लेंस सूत्र कौन-सा है?",
     ["1/v - 1/u = 1/f", "1/v + 1/u = 1/f", "v - u = f", "f = v + u"]),
    ("लेंस की क्षमता (P) और फोकस दूरी (f) में संबंध है — (f मीटर में)",
     ["P = f", "P = 1/f", "P = f²", "P = 1/f²"]),
    ("पूर्ण आंतरिक परावर्तन के लिए आवश्यक शर्त कौन-सी है?",
     ["प्रकाश विरल से सघन माध्यम में जाए", "आपतन कोण क्रांतिक कोण से बड़ा हो",
      "आपतन कोण क्रांतिक कोण से छोटा हो", "दोनों माध्यम समान हों"]),
    ("क्रांतिक कोण की स्थिति में अपवर्तन कोण का मान होता है —",
     ["0°", "45°", "90°", "180°"]),
    ("सूर्योदय से कुछ मिनट पहले सूर्य का दिखाई देना किस कारण होता है?",
     ["प्रकाश के परावर्तन", "प्रकाश के अपवर्तन", "प्रकाश के विक्षेपण", "प्रकाश के प्रकीर्णन"]),
    ("हीरे का चमकना किस घटना का उदाहरण है?",
     ["अपवर्तन", "विक्षेपण", "पूर्ण आंतरिक परावर्तन", "विवर्तन"]),
    ("1 डायोप्टर क्षमता वाले लेंस की फोकस दूरी होगी —",
     ["10 cm", "50 cm", "100 cm", "1000 cm"]),
    ("गोलीय लेंस में प्रकाशिक केंद्र से फोकस तक की दूरी कहलाती है —",
     ["वक्रता त्रिज्या", "फोकस दूरी", "द्वारक", "क्षमता"]),
]

LABELS = ["(A)", "(B)", "(C)", "(D)"]


def _ensure_typeface(rPr, tag, typeface):
    el = rPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        el.set("typeface", typeface)
        rPr.append(el)
    else:
        el.set("typeface", typeface)


def style_run(run, size, bold=False, italic=False, color=WHITE,
              latin=LATIN, cs=CSFONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor(*color)
    f.name = latin
    rPr = run._r.get_or_add_rPr()
    _ensure_typeface(rPr, "a:ea", latin)
    _ensure_typeface(rPr, "a:cs", cs)
    rPr.set("szCs", str(int(size * 100)))
    rPr.set("bCs", "1" if bold else "0")
    rPr.set("iCs", "1" if italic else "0")


def add_run(p, text, **kw):
    r = p.add_run()
    r.text = text
    style_run(r, **kw)
    return r


def new_para(tf, first=False, align=PP_ALIGN.LEFT, ls=1.0, before=0, after=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = ls
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    return p


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP, ml=0.05, mr=0.05, mt=0.02, mb=0.02):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(ml)
    tf.margin_right = Inches(mr)
    tf.margin_top = Inches(mt)
    tf.margin_bottom = Inches(mb)
    return tf


def rect_shape(slide, l, t, w, h, fill=None, line=None, line_w=1.0,
               shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor(*fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.fill.solid()
        sp.line.fill.fore_color.rgb = RGBColor(*line)
        sp.line.width = Pt(line_w)
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def paint_bg(slide):
    try:
        fill = slide.background.fill
        fill.gradient()
        fill.gradient_angle = 90.0
        stops = fill.gradient_stops
        stops[0].color.rgb = RGBColor(*BG_TOP)
        stops[0].position = 0.0
        stops[1].color.rgb = RGBColor(*BG_BOT)
        stops[1].position = 1.0
    except Exception:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*BG)


def q_size(text):
    n = len(text)
    if n <= 55:
        return 46
    if n <= 80:
        return 42
    if n <= 110:
        return 38
    if n <= 150:
        return 34
    return 32


def opt_size(opts):
    m = max(len(o) for o in opts)
    if m <= 16:
        return 30
    if m <= 26:
        return 28
    if m <= 40:
        return 25
    if m <= 60:
        return 23
    return 21


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "VVI Objective Questions - Refraction of Light - Dream Classes Kothwara"
prs.core_properties.author = "Dream Classes Kothwara"

# ================= TITLE SLIDE =================
s = prs.slides.add_slide(BLANK)
paint_bg(s)
rect_shape(s, 0, 0, 13.333, 0.06, fill=GOLD)  # top gold hairline

tf = textbox(s, 0.8, 1.05, 11.73, 0.75)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "विज्ञान (PHYSICS) — 10वीं कक्षा", size=32, bold=True, color=GOLD)

tf = textbox(s, 0.5, 1.85, 12.33, 1.5)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "VVI OBJECTIVE QUESTIONS", size=60, bold=True, color=WHITE)

rect_shape(s, 6.02, 3.5, 1.3, 0.045, fill=GOLD)

banner = rect_shape(s, 1.5, 3.8, 10.33, 1.2, fill=RED,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
btf = banner.text_frame
btf.word_wrap = True
btf.auto_size = None
btf.vertical_anchor = MSO_ANCHOR.MIDDLE
btf.margin_left = Inches(0.3); btf.margin_right = Inches(0.3)
p1 = new_para(btf, first=True, align=PP_ALIGN.CENTER, ls=1.0, after=4)
add_run(p1, "Chapter 2 — Refraction of Light", size=30, bold=True, color=WHITE)
p2 = new_para(btf, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p2, "प्रकाश का अपवर्तन", size=26, bold=True, color=GOLD)

tf = textbox(s, 0.8, 5.45, 11.73, 1.3)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0, after=6)
add_run(p, COACHING, size=30, bold=True, color=GOLD)
p = new_para(tf, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, f"By: {TEACHER}   |   Mob: {MOBILE}", size=20, bold=False, color=WHITE)

# ================= QUESTION SLIDES =================
for idx, (qtext, opts) in enumerate(QUESTIONS, start=1):
    s = prs.slides.add_slide(BLANK)
    paint_bg(s)

    # header strip
    rect_shape(s, 0, 0, 13.333, 0.62, fill=HEADER)
    rect_shape(s, 0, 0.62, 13.333, 0.03, fill=GOLD)
    tf = textbox(s, 0.5, 0.06, 9.8, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.LEFT, ls=1.0)
    add_run(p, SUBJECT_HEAD, size=20, bold=True, color=WHITE)

    badge = rect_shape(s, 11.05, 0.10, 1.78, 0.42, fill=GOLD,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
    btf = badge.text_frame
    btf.word_wrap = True; btf.auto_size = None
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    btf.margin_left = Inches(0.05); btf.margin_right = Inches(0.05)
    p = new_para(btf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, "10TH CLASS", size=14, bold=True, color=DARK)

    # red banner
    banner = rect_shape(s, 0.5, 0.85, 12.333, 0.62, fill=RED,
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    btf = banner.text_frame
    btf.word_wrap = True; btf.auto_size = None
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = new_para(btf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, "VVI OBJECTIVE QUESTION", size=24, bold=True, color=WHITE)
    rect_shape(s, 0.5, 1.47, 12.333, 0.035, fill=GOLD)

    # question number dot
    dot = rect_shape(s, 0.78, 0.94, 0.44, 0.44, fill=GOLD, shape=MSO_SHAPE.OVAL)
    dtf = dot.text_frame
    dtf.word_wrap = True; dtf.auto_size = None
    dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    dtf.margin_left = Inches(0); dtf.margin_right = Inches(0)
    p = new_para(dtf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, str(idx), size=20, bold=True, color=DARK)

    # question
    full_q = f"Q.{idx}. {qtext}"
    tf = textbox(s, 0.6, 1.68, 12.13, 2.0)
    p = new_para(tf, first=True, align=PP_ALIGN.LEFT, ls=1.06)
    add_run(p, full_q, size=q_size(full_q), bold=True, color=WHITE)

    # options 2x2
    osz = opt_size(opts)
    xs = [0.6, 6.79]
    ys = [3.88, 5.35]
    cw, ch = 5.94, 1.25
    for k, opt in enumerate(opts):
        card = rect_shape(s, xs[k % 2], ys[k // 2], cw, ch, fill=CARD,
                          line=GOLD, line_w=1.25,
                          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
        ctf = card.text_frame
        ctf.word_wrap = True; ctf.auto_size = None
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ctf.margin_left = Inches(0.22); ctf.margin_right = Inches(0.18)
        p = new_para(ctf, first=True, align=PP_ALIGN.LEFT, ls=1.04)
        add_run(p, LABELS[k] + "  ", size=osz, bold=True, color=GOLD)
        add_run(p, opt, size=osz, bold=False, color=WHITE)

    # footer
    rect_shape(s, 0.6, 6.82, 12.133, 0.015, fill=DIVIDER)
    tf = textbox(s, 0.6, 6.88, 5.5, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.LEFT, ls=1.0)
    add_run(p, "Dream Classes Kothwara", size=12, color=GREY)
    tf = textbox(s, 5.9, 6.88, 1.6, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
    add_run(p, MOBILE, size=11, color=GREY)
    tf = textbox(s, 10.7, 6.88, 2.03, 0.5)
    p = new_para(tf, first=True, align=PP_ALIGN.RIGHT, ls=1.0)
    add_run(p, TEACHER, size=12, color=GREY)

# ================= CLOSING SLIDE =================
s = prs.slides.add_slide(BLANK)
paint_bg(s)
rect_shape(s, 0, 0, 13.333, 0.06, fill=GOLD)

tf = textbox(s, 0.5, 1.5, 12.33, 1.5)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "THANK YOU", size=72, bold=True, color=WHITE)

tf = textbox(s, 0.5, 2.95, 12.33, 1.0)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "धन्यवाद!", size=44, bold=True, color=GOLD)

rect_shape(s, 6.02, 4.05, 1.3, 0.045, fill=GOLD)

tf = textbox(s, 0.8, 4.4, 11.73, 1.3)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0, after=6)
add_run(p, COACHING, size=30, bold=True, color=GOLD)
p = new_para(tf, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, f"By: {TEACHER}   |   Mob: {MOBILE}", size=20, color=WHITE)

tf = textbox(s, 0.8, 5.95, 11.73, 0.6)
p = new_para(tf, first=True, align=PP_ALIGN.CENTER, ls=1.0)
add_run(p, "Physics  •  प्रकाश का अपवर्तन  •  10वीं कक्षा", size=16, color=GREY)

out = "DREAMCLASSES_PHYSICS_Refraction-of-Light.pptx"
prs.save(out)
print(f"saved {out}: {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
